import openpyxl
from datetime import datetime
import re
import sys

def analyze_mass_balance(filename):
    """
    Analyzes the Mass Balance spreadsheet to find maximum shortage lengths per year.
    
    Args:
        filename: Path to the Excel file
        
    Returns:
        Dictionary mapping year to maximum consecutive shortage days
    """
    try:
        # Load the workbook
        wb = openpyxl.load_workbook(filename, data_only=True)
        
        # Get the "Daily Running Totals" sheet
        if "Daily Running Totals" not in wb.sheetnames:
            print("Error: 'Daily Running Totals' sheet not found in workbook")
            print(f"Available sheets: {wb.sheetnames}")
            return None
            
        sheet = wb["Daily Running Totals"]
        
        # Dictionary to store results
        max_shortages = {}
        
        # Read the header row to find year columns
        # Structure: Date, Rainfall (mm) YEAR, Surplus/Deficit YEAR,
        #            Running Total [Year], Shortages [Year], repeat...
        
        # Start from column B (column 2)
        col = 2
        years_found = []
        
        while True:
            header = sheet.cell(row=1, column=col).value
            if header is None:
                break
                
            # Check if this is a Rainfall column (start of a year's data)
            if header and "Rainfall (mm)" in str(header):
                # Extract the 4-digit year from the header. This works whether it is
                # written "Rainfall (mm) 2025" or "Rainfall (mm) [2025]".
                match = re.search(r"(\d{4})", str(header))
                if match:
                    years_found.append((int(match.group(1)), col))
            
            col += 1
        
        print(f"Found data for years: {[y[0] for y in years_found]}")
        
        # For each year, find the maximum value in the Shortages column
        for year, start_col in years_found:
            # Shortages column is 3 columns after Rainfall column
            # Rainfall=start_col, Surplus/Deficit=start_col+1, 
            # Running Total=start_col+2, Shortages=start_col+3
            shortages_col = start_col + 3
            
            max_shortage = 0

            # Scan the whole column (row 2 onward, after the header). Don't stop at
            # the first empty cell: non-leap years leave the Feb 29 row blank, and
            # stopping there would skip March through December.
            for row in range(2, sheet.max_row + 1):
                cell_value = sheet.cell(row=row, column=shortages_col).value

                if cell_value is None:
                    continue

                try:
                    shortage_value = float(cell_value)
                    max_shortage = max(max_shortage, shortage_value)
                except (ValueError, TypeError):
                    pass
            
            max_shortages[year] = max_shortage
        
        return max_shortages
        
    except FileNotFoundError:
        print(f"Error: File '{filename}' not found")
        return None
    except Exception as e:
        print(f"Error reading file: {e}")
        return None

def main():
    # Specify your file path, or pass one on the command line:
    #   python mass_balance.py "path/to/Mass Balance.xlsx"
    filename = sys.argv[1] if len(sys.argv) > 1 else "C:/Users/tag62/Downloads/Mass Balance.xlsx"
    
    print(f"Analyzing {filename}...")
    print("-" * 50)
    
    results = analyze_mass_balance(filename)
    
    if results:
        print("\nMaximum Shortage Length (days) by Year:")
        print("-" * 50)
        
        # Sort by year
        for year in sorted(results.keys()):
            max_days = results[year]
            print(f"{year}: {max_days:.0f} days")
        
        print("\n" + "=" * 50)
        print(f"Overall maximum shortage: {max(results.values()):.0f} days in year {max(results, key=results.get)}")

if __name__ == "__main__":
    main()